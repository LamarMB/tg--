"""Чтение/запись листов «Модули ПЛК» и «Каналы ПЛК»."""
from __future__ import annotations

from .model import INPUT_ELEMENTS, OUTPUT_ELEMENTS, PlcChannel, PlcModule

LAYOUT_KEYS = ("x0", "x1", "top", "h", "desc", "mark", "tip", "com_y", "com_dx", "pin", "coil",
               "bus", "ret", "mir", "ret_tip")


def layout_to_text(layout: dict) -> str:
    layout = layout or {}
    parts = [f"{k}={layout[k]:g}" for k in LAYOUT_KEYS if k in layout]
    parts += [f"{k}={v:g}" for k, v in (layout.get("xs") or {}).items()]
    return "; ".join(parts)


def layout_from_text(s: str) -> dict:
    out, xs = {}, {}
    for part in str(s or "").split(";"):
        if "=" in part:
            k, v = part.split("=", 1)
            k = k.strip()
            try:
                val = float(v.strip().replace(" ", "").replace(",", "."))
            except ValueError:
                continue
            (out if k in LAYOUT_KEYS else xs)[k] = val
    if xs:
        out["xs"] = xs
    return out


MOD_COLS = ["Модуль", "Позиция", "Тип", "Вид", "Ссылка", "Стрелка слева", "Стрелка вправо",
            "Марка провода", "Цвет", "Сечение", "Выходы NPN", "Раскладка"]
CH_COLS = ["Модуль", "Вывод", "Назначение", "Марка провода", "Цвет", "Сечение",
           "Элемент", "Обозначение", "Связь", "Ссылка", "Параметр"]

_EL_ALIASES = {"": "", "ттр": "ттр", "но ттр": "ттр", "перекл": "перекл",
               "перекидной": "перекл", "кнопка но": "кнопка но", "кнопка нз": "кнопка нз",
               "катушка": "катушка", "реле": "катушка", "лампа": "лампа",
               "общий": "общий", "com": "общий", "питание": "питание", "+24v": "питание"}


def _tag(s: str) -> str:
    s = s.strip()
    return s if not s or s.startswith("-") else "-" + s


def read_plc(wb, find_sheet, table, problems) -> list[PlcModule]:
    ws_m = find_sheet(wb, "Модули ПЛК", False, problems)
    ws_c = find_sheet(wb, "Каналы ПЛК", False, problems)
    if ws_m is None and ws_c is None:
        return []
    mods: dict[str, PlcModule] = {}
    order: list[PlcModule] = []
    if ws_m is not None:
        for rec in table(ws_m, MOD_COLS, problems):
            name = rec["Модуль"].lstrip("-")
            if not name:
                continue
            kind = rec["Вид"].lower()
            if kind.startswith("вх"):
                kind = "in"
            elif kind.startswith("вых"):
                kind = "out"
            else:
                problems.append(f"Лист «Модули ПЛК», модуль {name}: Вид «{rec['Вид']}» — "
                                "нужно «Входы» или «Выходы».")
                continue
            wire = [rec["Марка провода"], rec["Цвет"], rec["Сечение"]]
            m = PlcModule((rec["Позиция"] or name).lstrip("-"), rec["Тип"], kind, rec["Ссылка"],
                          npn=rec["Выходы NPN"].lower() in ("да", "1", "npn", "yes"))
            if kind == "in":
                m.feed, m.feed_next, m.feed_wire = rec["Стрелка слева"], rec["Стрелка вправо"], wire
            else:
                m.common, m.common_wire = rec["Стрелка слева"], wire
            key = name.upper()
            if key in mods:
                problems.append(f"Лист «Модули ПЛК»: модуль {name} указан дважды.")
                continue
            m.key = name
            m.layout = layout_from_text(rec.get("Раскладка", ""))
            mods[key] = m
            order.append(m)
    if ws_c is not None:
        last = ""
        for n, rec in enumerate(table(ws_c, CH_COLS, problems), 1):
            name = (rec["Модуль"] or last).lstrip("-")
            last = name
            m = mods.get(name.upper())
            where = f"Лист «Каналы ПЛК», модуль {name}, вывод {rec['Вывод'] or '?'}"
            if m is None:
                problems.append(f"{where}: модуль не описан на листе «Модули ПЛК».")
                continue
            el = _EL_ALIASES.get(rec["Элемент"].strip().lower())
            if el is None:
                problems.append(f"{where}: неизвестный элемент «{rec['Элемент']}».")
                continue
            allowed = INPUT_ELEMENTS if m.kind == "in" else OUTPUT_ELEMENTS
            if el not in allowed:
                problems.append(f"{where}: элемент «{rec['Элемент']}» не подходит для "
                                f"{'входов' if m.kind == 'in' else 'выходов'}.")
                continue
            if el in ("ттр", "перекл", "кнопка но", "кнопка нз", "катушка", "лампа") \
                    and not rec["Обозначение"]:
                problems.append(f"{where}: для элемента «{rec['Элемент']}» нужно Обозначение.")
                continue
            m.channels.append(PlcChannel(
                module=name, pin=rec["Вывод"], desc=rec["Назначение"],
                wire=[rec["Марка провода"], rec["Цвет"], rec["Сечение"]], element=el,
                device=_tag(rec["Обозначение"]), link=rec["Связь"], ref=rec["Ссылка"],
                param=rec["Параметр"]))
    return [m for m in order if m.channels]


def write_plc(wb, modules: list[PlcModule], header) -> None:
    ws = wb.create_sheet("Модули ПЛК")
    header(ws, MOD_COLS, [9, 9, 18, 9, 8, 22, 22, 13, 8, 8, 10, 30])
    for m in modules:
        wire = m.feed_wire if m.kind == "in" else m.common_wire
        wire = (wire + ["", "", ""])[:3]
        ws.append([m.key or m.tag, m.tag if m.key and m.key != m.tag else "", m.type, "Входы" if m.kind == "in" else "Выходы", m.ref,
                   m.feed if m.kind == "in" else m.common,
                   m.feed_next if m.kind == "in" else "", *wire,
                   "да" if m.npn else "", layout_to_text(m.layout)])
    ws2 = wb.create_sheet("Каналы ПЛК")
    header(ws2, CH_COLS, [8, 7, 34, 12, 7, 8, 11, 11, 18, 8, 9])
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", allow_blank=True, formula1='"ТТР,Перекл,Кнопка НО,'
                        'Кнопка НЗ,Катушка,Лампа,Общий,Питание"')
    ws2.add_data_validation(dv)
    dv.add("G2:G3000")
    names = {"ттр": "ТТР", "перекл": "Перекл", "кнопка но": "Кнопка НО",
             "кнопка нз": "Кнопка НЗ", "катушка": "Катушка", "лампа": "Лампа",
             "общий": "Общий", "питание": "Питание", "": ""}
    for m in modules:
        for c in m.channels:
            w = (c.wire + ["", "", ""])[:3]
            ws2.append([m.key or m.tag, c.pin, c.desc, *w, names[c.element], c.device, c.link,
                        c.ref, c.param])


# ------------------------------------------------------------ Питание 24 В
P24_COLS = ["Группа", "Тип", "Обозначение", "Номинал", "Марка провода", "Цвет",
            "Сечение", "Связь", "Ссылка", "Связь 2", "Ссылка 2", "Подпись", "Раскладка",
            "Провод источника"]
P24_TYPES = {"ввод": "ввод", "источник": "ввод", "автомат": "автомат", "qf": "автомат",
             "отвод вверх": "вверх", "вверх": "вверх", "отвод вниз": "вниз", "вниз": "вниз",
             "клемма": "клемма", "клемма ввода": "клемма"}


def _split_wire(s: str) -> list[str]:
    """«L+ / RD / 1,5» → [L+, RD, 1,5]."""
    parts = [x.strip() for x in str(s or "").replace(";", "/").split("/")]
    return parts if any(parts) else []


def read_power24(wb, find_sheet, table, problems):
    from .power24 import (Breaker, BreakerGroup, InClamp, InputBlock, MinusBus, MinusTap,
                          Power24)
    ws = find_sheet(wb, "Питание 24В", False, problems) or \
        find_sheet(wb, "Питание 24 В", False, problems)
    if ws is None:
        return None
    rows, last = [], ""
    for rec in table(ws, P24_COLS, problems):
        grp = rec["Группа"] or last
        last = grp
        kind = P24_TYPES.get(rec["Тип"].strip().lower())
        if not grp or kind is None:
            problems.append(f"Лист «Питание 24В», группа {grp or '?'}: тип «{rec['Тип']}» — "
                            "нужно: ввод, автомат, отвод вверх, отвод вниз, клемма.")
            continue
        rows.append((grp, kind, rec))
    pw = Power24()
    groups: dict[str, object] = {}
    blocks: dict[str, InputBlock] = {}
    for grp, kind, rec in rows:
        if kind == "клемма":                       # вводной клеммник (-X0.3)
            blk = blocks.get(grp)
            if blk is None:
                blk = blocks[grp] = InputBlock(grp.lstrip("-"))
                pw.inputs.append(blk)
            lay = layout_from_text(rec.get("Раскладка", ""))
            blk.clamps.append(InClamp(
                rec["Обозначение"], rec["Связь"], rec["Ссылка"],
                _split_wire(rec.get("Провод источника", "")),
                [rec["Марка провода"], rec["Цвет"], rec["Сечение"]],
                rec["Связь 2"], rec["Ссылка 2"], float(lay.get("x0", 0) or 0)))
            if rec["Подпись"]:
                blk.note = rec["Подпись"]
            continue
        is_minus = any(g == grp and k in ("вверх", "вниз") for g, k, _ in rows)
        obj = groups.get(grp)
        if obj is None:
            obj = MinusBus(grp.lstrip("-")) if is_minus else BreakerGroup(grp)
            groups[grp] = obj
            (pw.minus if is_minus else pw.groups).append(obj)
        wire = [rec["Марка провода"], rec["Цвет"], rec["Сечение"]]
        if is_minus and rec.get("Раскладка"):
            obj.layout = layout_from_text(rec["Раскладка"])
        if kind == "ввод":
            obj.source, obj.source_ref, obj.source_wire = rec["Связь"], rec["Ссылка"], wire
        elif kind == "автомат":
            if is_minus:
                problems.append(f"Лист «Питание 24В», группа {grp}: автоматы и отводы "
                                "минуса в одной группе.")
                continue
            tag = rec["Обозначение"]
            tag = tag if tag.startswith("-") or not tag else "-" + tag
            targets = [(rec["Связь"], rec["Ссылка"])]
            if rec["Связь 2"]:
                targets.append((rec["Связь 2"], rec["Ссылка 2"]))
            obj.breakers.append(Breaker(tag, rec["Номинал"], wire, targets, rec["Подпись"]))
        else:
            obj.taps.append(MinusTap(rec["Обозначение"], kind == "вверх", wire,
                                     rec["Связь"], rec["Ссылка"]))
    return pw if pw else None


def write_power24(wb, pw, header) -> None:
    ws = wb.create_sheet("Питание 24В")
    header(ws, P24_COLS, [8, 11, 12, 12, 12, 8, 8, 20, 8, 16, 8, 24, 24, 16])
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", allow_blank=True,
                        formula1='"ввод,автомат,отвод вверх,отвод вниз,клемма"')
    ws.add_data_validation(dv)
    dv.add("B2:B2000")
    if not pw:
        return
    for g in pw.groups:
        if g.source:
            ws.append([g.name, "ввод", "", "", *(g.source_wire + ["", "", ""])[:3],
                       g.source, g.source_ref, "", "", ""])
        for b in g.breakers:
            t = b.targets + [("", "")] * (2 - len(b.targets))
            ws.append([g.name, "автомат", b.tag, b.rating, *(b.wire + ["", "", ""])[:3],
                       t[0][0], t[0][1], t[1][0], t[1][1], b.caption])
    for b in pw.inputs:
        note = b.note
        for c in b.clamps:
            ws.append([b.name, "клемма", c.name, "", *(c.wire + ["", "", ""])[:3],
                       c.source, c.source_ref, c.feed, c.jumper, note,
                       f"x0={c.x:g}" if c.x else "", " / ".join(c.source_wire)])
            note = ""
    for m in pw.minus:
        lay = layout_to_text(m.layout)
        if m.source:
            ws.append([m.name, "ввод", "", "", *(m.source_wire + ["", "", ""])[:3],
                       m.source, m.source_ref, "", "", "", lay])
            lay = ""
        for t in m.taps:
            ws.append([m.name, "отвод вверх" if t.up else "отвод вниз", t.clamp, "",
                       *(t.wire + ["", "", ""])[:3], t.link, t.ref, "", "", "", lay])
            lay = ""


# ------------------------------------------------------------ Линии 230 В
F_COLS = ["Автомат", "Тип", "Номинал", "Утечка", "Сечение", "Питание L", "Ссылка L",
          "Питание N", "Ссылка N", "Клеммник", "Кабель", "Марка кабеля", "Жилы", "Зона",
          "Розетки", "Номинал розеток", "Нагрузка L", "Нагрузка N", "Нагрузка PE",
          "Ссылка нагрузки", "Подпись"]


def read_feeders(wb, find_sheet, table, problems):
    from .power230 import Feeder
    ws = find_sheet(wb, "Линии 230В", False, problems) or \
        find_sheet(wb, "Линии 230 В", False, problems)
    if ws is None:
        return []
    out = []
    for rec in table(ws, F_COLS, problems):
        tag = rec["Автомат"]
        if not tag:
            problems.append("Лист «Линии 230В»: строка без обозначения автомата.")
            continue
        kind = rec["Тип"].strip().lower() or "авдт"
        if kind not in ("авдт", "ав"):
            problems.append(f"Лист «Линии 230В», {tag}: тип «{rec['Тип']}» — нужно АВДТ или АВ.")
            continue
        socks = [x.strip() for x in rec["Розетки"].replace(",", ";").split(";") if x.strip()]
        loads = [rec["Нагрузка L"], rec["Нагрузка N"], rec["Нагрузка PE"]]
        out.append(Feeder(tag if tag.startswith("-") else "-" + tag, kind, rec["Номинал"],
                          rec["Утечка"], rec["Сечение"] or "2,5", rec["Питание L"],
                          rec["Ссылка L"], rec["Питание N"], rec["Ссылка N"],
                          rec["Клеммник"], rec["Кабель"], rec["Марка кабеля"], rec["Жилы"],
                          rec["Зона"], socks, rec["Номинал розеток"],
                          loads if any(loads) else [], rec["Ссылка нагрузки"], rec["Подпись"]))
    return out


def write_feeders(wb, feeders, header) -> None:
    ws = wb.create_sheet("Линии 230В")
    header(ws, F_COLS, [9, 7, 9, 8, 8, 12, 8, 12, 8, 9, 9, 24, 8, 7, 18, 9, 14, 14, 14, 10, 30])
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", allow_blank=True, formula1='"АВДТ,АВ"')
    ws.add_data_validation(dv)
    dv.add("B2:B500")
    for f in feeders or []:
        loads = (f.load_links + ["", "", ""])[:3]
        ws.append([f.tag, "АВДТ" if f.kind == "авдт" else "АВ", f.rating, f.leak, f.section,
                   f.source, f.source_ref, f.n_source, f.n_ref, f.terminal, f.cable,
                   f.cable_type, f.cable_cores, f.zone, ";".join(f.sockets), f.socket_rating,
                   *loads, f.load_ref, f.caption])


# ------------------------------------------------------------ Ввод 230 В
M_COLS = ["Тип", "Обозначение", "Номинал", "Название", "Вывод", "Сверху",
          "Нагрузка", "Нагрузка обозн.", "Нагрузка парам.", "Нагрузка 2 обозн.",
          "Нагрузка 2 парам.", "Марка провода", "Цвет", "Сечение", "Связь", "Ссылка",
          "Нейтраль"]
M_TYPES = ("ввод", "выключатель", "выход", "нейтраль", "отвод n", "ветвь", "устройство",
           "вывод")


def _t(s: str) -> str:
    s = (s or "").strip()
    return s if not s or s.startswith(("-", "+")) else "-" + s


def read_mains(wb, find_sheet, table, problems):
    from .mains import Branch, Device, Mains, NTap, Pin, QsOutput
    ws = find_sheet(wb, "Ввод 230В", False, problems) or \
        find_sheet(wb, "Ввод 230 В", False, problems)
    if ws is None:
        return None
    m = Mains()
    devices: dict[str, Device] = {}
    last_dev = None
    for rec in table(ws, M_COLS, problems):
        kind = rec["Тип"].strip().lower()
        wire = [rec["Марка провода"], rec["Цвет"], rec["Сечение"]]
        if kind == "ввод":
            m.input_block = rec["Обозначение"].lstrip("-") or m.input_block
            if rec["Вывод"]:
                m.input_labels = rec["Вывод"].replace(",", " ").replace(";", " ").split()
            m.input_from = rec["Связь"]
            m.input_cable = rec["Название"]
            m.input_cable_type = rec["Номинал"]
        elif kind == "выключатель":
            m.qs = rec["Обозначение"].lstrip("-") or m.qs
            m.qs_rating = rec["Номинал"]
            if rec["Вывод"].isdigit():
                m.qs_poles = max(1, min(6, int(rec["Вывод"])))
        elif kind == "выход":
            m.qs_outputs.append(QsOutput(rec["Вывод"], rec["Связь"], rec["Ссылка"], wire))
        elif kind == "нейтраль":
            m.n_bus = rec["Обозначение"].lstrip("-") or m.n_bus
        elif kind == "отвод n":
            m.n_taps.append(NTap(rec["Вывод"], rec["Связь"], rec["Ссылка"], wire))
        elif kind == "ветвь":
            m.branches.append(Branch(
                _t(rec["Обозначение"]), rec["Номинал"], wire,
                rec["Нагрузка"].strip().lower() or "стрелка", _t(rec["Нагрузка обозн."]),
                rec["Нагрузка парам."], _t(rec["Нагрузка 2 обозн."]), rec["Нагрузка 2 парам."],
                rec["Нейтраль"], rec["Связь"], rec["Ссылка"]))
        elif kind == "устройство":
            tag = _t(rec["Обозначение"])
            last_dev = devices[tag.upper()] = Device(tag, rec["Название"], rec["Номинал"])
        elif kind == "вывод":
            tag = _t(rec["Обозначение"])
            dev = devices.get(tag.upper()) if tag else last_dev
            if dev is None:
                problems.append(f"Лист «Ввод 230В»: вывод {rec['Вывод']} — устройство "
                                f"«{tag}» не описано выше строкой «устройство».")
                continue
            top = rec["Сверху"].strip().lower() not in ("нет", "низ", "снизу", "0", "false")
            dev.pins.append(Pin(rec["Вывод"], top, rec["Связь"], rec["Ссылка"], wire))
        else:
            problems.append(f"Лист «Ввод 230В»: тип «{rec['Тип']}» — нужно: "
                            + ", ".join(M_TYPES) + ".")
    # устройство, которое питается от ветви, рисуется в ней
    for b in m.branches:
        if b.load.startswith("устр"):
            b.device = devices.pop(b.load_tag.upper(), None)
    m.devices = list(devices.values())
    return m if m else None


def write_mains(wb, m, header) -> None:
    ws = wb.create_sheet("Ввод 230В")
    header(ws, M_COLS, [11, 12, 22, 24, 8, 7, 11, 11, 14, 11, 16, 11, 7, 7, 14, 8, 10])
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", allow_blank=True, formula1='"' + ",".join(M_TYPES) + '"')
    ws.add_data_validation(dv)
    dv.add("A2:A1000")
    dv2 = DataValidation(type="list", allow_blank=True,
                         formula1='"розетка,лампа,термостат,светильник,устройство,стрелка"')
    ws.add_data_validation(dv2)
    dv2.add("G2:G1000")
    if not m:
        return
    E = [""] * len(M_COLS)

    def row(**kv):
        r = list(E)
        for k, v in kv.items():
            r[M_COLS.index(k)] = v
        ws.append(r)

    def w3(w):
        w = (list(w or []) + ["", "", ""])[:3]
        return {"Марка провода": w[0], "Цвет": w[1], "Сечение": w[2]}

    def dev(d):
        row(**{"Тип": "устройство", "Обозначение": d.tag, "Номинал": d.param,
               "Название": d.title})
        for pn in d.pins:
            row(**{"Тип": "вывод", "Обозначение": d.tag, "Вывод": pn.name,
                   "Сверху": "да" if pn.top else "нет", "Связь": pn.link, "Ссылка": pn.ref,
                   **w3(pn.wire)})

    row(**{"Тип": "ввод", "Обозначение": m.input_block, "Вывод": " ".join(m.input_labels),
           "Связь": m.input_from, "Название": m.input_cable, "Номинал": m.input_cable_type})
    row(**{"Тип": "выключатель", "Обозначение": m.qs, "Номинал": m.qs_rating,
           "Вывод": str(m.qs_poles)})
    for o in m.qs_outputs:
        row(**{"Тип": "выход", "Вывод": o.pole, "Связь": o.link, "Ссылка": o.ref,
               **w3(o.wire)})
    row(**{"Тип": "нейтраль", "Обозначение": m.n_bus})
    for t in m.n_taps:
        row(**{"Тип": "отвод n", "Вывод": t.clamp, "Связь": t.link, "Ссылка": t.ref,
               **w3(t.wire)})
    for b in m.branches:
        row(**{"Тип": "ветвь", "Обозначение": b.tag, "Номинал": b.rating, "Нагрузка": b.load,
               "Нагрузка обозн.": b.load_tag or (b.device.tag if b.device else ""),
               "Нагрузка парам.": b.load_param, "Нагрузка 2 обозн.": b.load_tag2,
               "Нагрузка 2 парам.": b.load_param2, "Нейтраль": b.n_link,
               "Связь": b.link, "Ссылка": b.ref, **w3(b.wire)})
        if b.device:
            dev(b.device)
    for d in m.devices:
        dev(d)


# ------------------------------------------------------------ Сеть
N_COLS = ["Тип", "Устройство", "Имя", "Вид", "Сторона", "Связь", "Ссылка",
          "Марка провода", "Цвет", "Сечение"]
L_COLS = ["Кабель", "Марка", "Длина", "Откуда", "Куда", "Разъём", "Вид разъёма", "Зона",
          "Удалённый конец", "Ссылка", "Кабель дальше", "Марка дальше", "Устройство",
          "Название устройства", "Порт устройства", "Примечание"]


def read_network(wb, find_sheet, table, problems):
    from .network import NetDevice, NetLink, NetPort, NetPower, Network
    ws = find_sheet(wb, "Сеть", False, problems)
    wl = find_sheet(wb, "Кабели сети", False, problems)
    if ws is None and wl is None:
        return None
    net = Network()
    devs: dict[str, NetDevice] = {}
    if ws is not None:
        for rec in table(ws, N_COLS, problems):
            kind = rec["Тип"].strip().lower()
            tag = _t(rec["Устройство"])
            if kind == "устройство":
                d = devs[tag.upper()] = NetDevice(tag, rec["Имя"], rec["Вид"].strip().lower(),
                                                  rec["Сторона"].strip().lower() or "top",
                                                  rec["Ссылка"], pe=rec["Связь"].upper() == "PE")
                net.devices.append(d)
                continue
            d = devs.get(tag.upper())
            if d is None:
                problems.append(f"Лист «Сеть»: «{tag}» не описано строкой «устройство».")
                continue
            if kind == "порт":
                d.ports.append(NetPort(rec["Имя"], rec["Вид"].strip().lower() or "rj45",
                                       rec["Сторона"].strip().lower() or "top"))
            elif kind == "питание":
                d.power.append(NetPower(rec["Имя"], rec["Связь"], rec["Ссылка"],
                                        [rec["Марка провода"], rec["Цвет"], rec["Сечение"]]))
            elif kind == "модуль":
                d.modules.append(rec["Имя"])
            else:
                problems.append(f"Лист «Сеть»: тип «{rec['Тип']}» — нужно: устройство, порт, "
                                "питание, модуль.")
    if wl is not None:
        for r in table(wl, L_COLS, problems):
            net.links.append(NetLink(r["Кабель"], r["Марка"], r["Длина"], r["Откуда"], r["Куда"],
                                     r["Разъём"], r["Вид разъёма"] or "RJ45", r["Зона"],
                                     r["Удалённый конец"], r["Ссылка"], r["Кабель дальше"],
                                     r["Марка дальше"], r["Устройство"],
                                     r["Название устройства"].replace("\\n", "\n"),
                                     r["Порт устройства"], r["Примечание"].replace("\\n", "\n")))
    return net if net else None


def write_network(wb, net, header) -> None:
    ws = wb.create_sheet("Сеть")
    header(ws, N_COLS, [12, 12, 36, 10, 9, 14, 8, 11, 7, 7])
    wl = wb.create_sheet("Кабели сети")
    header(wl, L_COLS, [9, 18, 8, 16, 16, 9, 10, 8, 20, 10, 12, 12, 11, 20, 10, 24])
    if not net:
        return
    for d in net.devices:
        ws.append(["устройство", d.tag, d.name, d.brand, d.row, "PE" if d.pe else "", d.ref,
                   "", "", ""])
        for p in d.ports:
            ws.append(["порт", d.tag, p.name, p.kind, p.side, "", "", "", "", ""])
        for pw in d.power:
            w = (list(pw.wire) + ["", "", ""])[:3]
            ws.append(["питание", d.tag, pw.pin, "", "", pw.link, pw.ref, *w])
        for m in d.modules:
            ws.append(["модуль", d.tag, m, "", "", "", "", "", "", ""])
    for l in net.links:
        wl.append([l.cable, l.cable_type, l.length, l.a, l.b, l.socket, l.socket_kind, l.zone,
                   l.remote, l.remote_ref, l.remote_cable, l.remote_cable_type, l.target,
                   l.target_title.replace("\n", "\\n"), l.target_port,
                   l.note.replace("\n", "\\n")])


# ------------------------------------------------------------ Коробки / внешние шкафы
F2_COLS = ["Зона", "Лист", "Клеммник", "Кабель", "Ссылка кабеля", "Марка кабеля", "Экран",
           "PE", "Парами", "Удалённое устройство", "Клемма", "Жила", "Удалённый конец",
           "Ссылка конца", "Направление", "Марка провода", "Цвет", "Сечение", "Элемент",
           "Обозначение", "Параметр", "Связь", "Ссылка", "Подпись", "Перемычка"]


def _yes(v: str) -> bool:
    return str(v or "").strip().lower() in ("да", "yes", "1", "true", "+")


def read_fields(wb, find_sheet, table, problems):
    from .field import FieldArea, FieldGroup, FieldTerm
    ws = find_sheet(wb, "Коробки", False, problems)
    if ws is None:
        return []
    areas: list[FieldArea] = []
    area = grp = None
    for r in table(ws, F2_COLS, problems):
        zone = r["Зона"] or (area.zone if area else "")
        sheet = r["Лист"]
        if area is None or r["Зона"] and (zone != area.zone or (sheet and sheet != area.sheet)):
            area = FieldArea(zone, [], sheet)
            areas.append(area)
            grp = None
        block = r["Клеммник"] or (grp.block if grp else "")
        if not block:
            problems.append("Лист «Коробки»: строка без клеммника.")
            continue
        if grp is None or block != grp.block:
            grp = FieldGroup(block, r["Кабель"], r["Ссылка кабеля"], r["Марка кабеля"],
                             _yes(r["Экран"]), _yes(r["PE"]), _yes(r["Парами"]),
                             r["Удалённое устройство"])
            area.groups.append(grp)
        if not r["Клемма"]:
            continue
        grp.terms.append(FieldTerm(
            r["Клемма"], r["Жила"], r["Удалённый конец"], r["Ссылка конца"],
            "out" if r["Направление"].strip().lower() in ("out", "в поле", "из шкафа") else "in",
            [r["Марка провода"], r["Цвет"], r["Сечение"]], r["Элемент"].strip().lower(),
            _t(r["Обозначение"]), r["Параметр"].replace("\\n", "\n"), r["Связь"], r["Ссылка"],
            r["Подпись"].replace("\\n", "\n"), r["Перемычка"]))
    return areas


def write_fields(wb, areas, header) -> None:
    ws = wb.create_sheet("Коробки")
    header(ws, F2_COLS, [7, 5, 9, 11, 10, 26, 6, 5, 7, 10, 7, 5, 18, 10, 9, 11, 7, 7, 9, 9,
                         12, 12, 7, 28, 9])
    for a in areas or []:
        first_area = True
        for g in a.groups:
            first = True
            for t in g.terms or [None]:
                head = [a.zone if first_area else "", a.sheet if first_area else "",
                        g.block, g.cable if first else "", g.cable_ref if first else "",
                        g.cable_type if first else "", ("да" if g.shield else "") if first else "",
                        ("да" if g.pe else "") if first else "", ("да" if g.paired else "") if first else "",
                        g.remote_device if first else ""]
                if t is None:
                    ws.append(head + [""] * 15)
                else:
                    w = (list(t.wire) + ["", "", ""])[:3]
                    ws.append(head + [t.clamp, t.core, t.remote, t.remote_ref, t.dir, *w, t.element,
                                      t.device, t.param.replace("\n", "\\n"), t.link, t.ref,
                                      t.caption.replace("\n", "\\n"), t.bridge])
                first = first_area = False


# ------------------------------------------------------------ Сигнальная колонна
C_COLS = ["Тип", "Вывод", "Цвет", "Вид", "Реле", "Марка провода", "Цвет провода", "Сечение",
          "Провод от реле", "Через контакт", "Выводы контакта", "Ссылка", "Связь"]


def read_column(wb, find_sheet, table, problems):
    from .column import ColLamp, Column
    ws = find_sheet(wb, "Колонна", False, problems)
    if ws is None:
        return None
    col = Column()
    for r in table(ws, C_COLS, problems):
        kind = r["Тип"].strip().lower()
        wire = [r["Марка провода"], r["Цвет провода"], r["Сечение"]]
        if kind == "колонна":
            col.tag, col.title = _t(r["Реле"]) or col.tag, r["Цвет"] or col.title
        elif kind == "питание":
            col.feed_wire, col.feed_next, col.feed_next_ref = wire, r["Связь"], r["Ссылка"]
        elif kind == "общий":
            col.common_pin, col.common_link, col.common_ref = r["Вывод"] or "0", r["Связь"], r["Ссылка"]
            col.common_wire = wire
        elif kind == "сегмент":
            rw = [x.strip() for x in r["Провод от реле"].split("/")] if r["Провод от реле"] else []
            col.lamps.append(ColLamp(r["Вывод"], r["Цвет"], r["Вид"] or "лампа", _t(r["Реле"]),
                                     wire, rw, _t(r["Через контакт"]),
                                     r["Выводы контакта"] or "11/14/12", r["Ссылка"]))
    return col if col else None


def write_column(wb, col, header) -> None:
    if not col:
        return
    ws = wb.create_sheet("Колонна")
    header(ws, C_COLS, [10, 7, 18, 9, 8, 12, 10, 8, 18, 12, 12, 8, 14])
    fw = (list(col.feed_wire) + ["", "", ""])[:3]
    cw = (list(col.common_wire) + ["", "", ""])[:3]
    ws.append(["колонна", "", col.title, "", col.tag, "", "", "", "", "", "", "", ""])
    ws.append(["питание", "", "", "", "", *fw, "", "", "", col.feed_next_ref, col.feed_next])
    for l in col.lamps:
        w = (list(l.wire) + ["", "", ""])[:3]
        ws.append(["сегмент", l.pin, l.color, l.kind, l.relay, *w, " / ".join(l.relay_wire),
                   l.via, l.via_pins if l.via else "", l.via_ref, ""])
    ws.append(["общий", col.common_pin, "", "", "", *cw, "", "", "", col.common_ref, col.common_link])
